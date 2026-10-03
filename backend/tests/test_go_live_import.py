"""Go-live data load: cleaning MADIO's hand-kept sheets, preview, the
archived wipe-and-load (company-scoped), undo, and the starter flows."""
import asyncio
import io
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from fastapi import HTTPException
from mongomock_motor import AsyncMongoMockClient
from starlette.datastructures import UploadFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import go_live_import as gl  # noqa: E402
import server  # noqa: E402

ADMIN = {"id": "u1", "tenant_id": "madio", "name": "Admin", "role": "admin", "username": "admin"}
OTHER = {"id": "u9", "tenant_id": "studio", "name": "S", "role": "admin", "username": "s"}


@pytest.fixture(autouse=True)
def _db(monkeypatch):
    monkeypatch.setattr(server, "db", AsyncMongoMockClient()["go_live"])


def run(c):
    return asyncio.run(c)


def _book(sheets: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ENQUIRY = {
    "Visitors": [
        (None, "MADIO FURNITURE"),
        ("Sl. No", "Date", "Cust Name & Location", "Reference", "Phone Number", "Requirement",
         "Attend person", "Site Visit (Date/Exe)", "Remarks  ", "STATUS & BALANCE", "Ticket Value"),
        (1, datetime(2026, 4, 1), "Kiran - Gachibowli", "Walk In", 9876500001.0, "Sofa", "Rooth", None,
         "Need to share quotation", None, None),
        (2, "15/09/25", "Meena Rao, Shaikpet", "Raghu MF", "+91 98765 00002", "Paint texture", "GK", None,
         "Deal closed", "Adv Received", 60000),
        (3, None, "Kishore - Hyd", "walkin", "Not given", "Furniture", "Rooth", None, "", "Cancel", None),
        (4, None, None, None, None, None, None, None, None, None, None),
    ],
    "2026 MF Quotes": [
        ("Sl. No", "Quot. No.", "Client Visit Date", "Cust Name & Location", "Reference", "Phone Number",
         "Requirement", "NHPC Q.NO ", "Quot. Date", "Quot. Through", "Attend person", "Site visit Date / Exe",
         "Remarks  ", "STATUS & BALANCE", "Ticket Value", "Cash", "Bank"),
        (1, "AF - 2601-001", "04//01/26", "Kiran - Gachibowli", "Walk In", 9876500001.0, "Furniture", None,
         datetime(2026, 4, 1), "Whatsapp", "Rooth", None, "Amount Received (4/01)", "Delivery Done",
         171168.0, 56168.0, 11378.0),
        (2, "AF-2601-002", "?", "Dr Arun - Kohinoor", "Walk In", "-", "Doors&Windows", None, None, "Whatsapp",
         "Jags", None, "Low Budget", None, None, None, None),
        (3, "AF-2601-003", "05/01/26", "Ravi - Kokapet", "Ar Rekha", 9876543210.0, "MAP", None, "05/01/26",
         "Whatsapp", "GK", None, "Shared quotation, follow up", None, "1,20,000", None, None),
        (4, None, None, "Nameless row with no number", None, None, None, None, None, None, None, None, None,
         None, None, None, None),
    ],
    "MF Sale": [
        ("Sl. No", "Sale No", "Adv / Cof Date", "Cust Name & Location", "Reference", "Phone Number",
         "Requirement", "Quot. Date", "Quote No", "Attend person", "Site visit Date / Exe", "Remarks  ",
         "STATUS & BALANCE", "Ticket Value & Mode", None, "Bank", "cash", "2nd", "3rd"),
        (1, "MF 001", "06/01/26", "Kiran - Gachibowli", "Walk In", 9876500001.0, "Sofa", None,
         "AF-2601-001", "Rooth", None, "Amount received", "Delivered", "1,71,168", "x", 11378.0, 56168.0,
         "-", "-"),
        (2, "MF 002", "20/09/26", "Walk-in buyer - Hyd", "Walk In", None, "Chairs", None, "", "Rooth", None,
         "Advance", 30000.0, 50000.0, None, "-", 20000.0, None, None),
        (None, None, None, "Repeated enquiry tail", None, None, None, None, None, None, None, None, None,
         None, None, None, None, None, None),
    ],
    "Arch": [
        (1, "Ar Prakash - PK Studio, Yousafguda", "98765 00003", "Arch", "Gift", "Y"),
        (2, "Latha M - Colour Works", "98765 00004", "Interior", "Gift", None, "Y"),
    ],
}
PURCHASES = {
    "Purchase Order": [
        (66080, "DATE", "PO  - NUMBER", "Quotation Number", "CUSTOMER / STOCK", "VENDOR", "PRODUCT DETAILS",
         "QTY", "VALUE", "INVOICE NO", "INVOICE DATE", "Receive Date"),
        (1, "30-4-24", "AF/24-25/007", None, None, "Dextra Square Pvt LTD", None, "788.90 sft", "1,22,488",
         "TS/5/24-25", "30-4-24", None),
        (2, "21-05-25", "AF/25-26/82", "AF-2601-003/2", "Ravi", "HEURISTIC FORMULATION", "Silcol 05 kg",
         "3 cont", 7025.0, "25-26/AHF/814", "22-5-25", "23-5-25"),
        (3, None, None, None, None, None, "Batch : 250612AEDC02", None, None, None, None, None),
        (4, None, None, None, None, None, "Decor Matt 500 gm", "1 cont", 500.0, None, None, None),
        (5, "22-05-25", "AF/25-26/82", None, None, None, "Pearl Polo", "2 cont", 1000.0, None, None, None),
        (6, "25-05-25", "AF/25-26/83", None, "Shop", None, "Cimento", "1 cont", 900.0, "25-26/AHF/900", None,
         None),
    ],
    "Sheet2": [
        (None, "MAP STOCK LIST"),
        (None, "sl.No", "Po Number", "Product Name", "Cust", "Batch Number", "Inflow", "Out flow", "Qty", "Cont"),
        (None, 1, None, "Cimento BD", None, None, None, None, "5kg", 3),
    ],
}
MIS = {
    "Sales Accounts": [("Madio Furniture",), ("PLOT NO 25 , ROAD NO 1",), ("KONDAPUR , HYDERABAD -500084",),
                       ("GST IN - 36ACDFA4831E1ZR",), ("E-Mail : LUXURY.ARMADIO@GMAIL.COM",)],
    "Closing Stock": [
        (None, None, None, None, None, None, None, None, None, None, 1, 2, None),
        ("Month", "Purchase Date", "PRODUCT", "VENDOR", "MODEL NO", "QTY", "RATE", "LP", "Taxable Value",
         "GST - 18%", "Purchase Cost", "Selling Price", "Status"),
        (datetime(2026, 4, 1), datetime(2024, 3, 31), "Ficus Microcarpa", "  V1- ZHI RAN SHE",
         "Artificial Plants", 1, 32000, 32000, 27119, 4881, 32000, 98000, "Display"),
        (datetime(2026, 4, 1), datetime(2026, 4, 21), "10115021", "Dining Chair", "Novanthe Mobila Pvt Ltd",
         4, 5500, 22000, 18644, 3355, 22000, 66000, "Godown   "),
    ],
}


def _files():
    return [UploadFile(io.BytesIO(_book(ENQUIRY)), filename="AF_Sheet.xlsx"),
            UploadFile(io.BytesIO(_book(PURCHASES)), filename="Purchase_Order.xlsx"),
            UploadFile(io.BytesIO(_book(MIS)), filename="MIS.xlsx")]


# ── cleaning ───────────────────────────────────────────────────────────────
def test_cell_cleaning():
    # Excel turned day/month into month/day for any day <= 12; swapped back.
    assert gl.sheet_date(datetime(2026, 4, 1)) == "2026-01-04"
    assert gl.sheet_date(datetime(2026, 4, 21)) == "2026-04-21"
    assert gl.sheet_date("04//01/26") == "2026-01-04"
    assert gl.sheet_date("15/09/25") == "2025-09-15"
    assert gl.sheet_date("?") == "" and gl.sheet_date(45000) == ""
    assert gl.phone(9876500001.0) == "9876500001"
    assert gl.phone("+91 98765 00002") == "9876500002"
    assert gl.phone("919246531773") == "9246531773"
    assert gl.phone("Not given") == "" and gl.phone("-") == ""
    assert gl.amount("2,60,000") == 260000 and gl.amount("50,000 B.T") == 50000
    assert gl.amount("2,00,000/-") == 200000 and gl.amount("-") == 0
    assert gl.name_location("Meena Rao, Shaikpet") == ("Meena Rao", "Shaikpet")
    assert gl.name_location("Sudheer - Jubleehills") == ("Sudheer", "Jubleehills")
    assert gl.name_location("Pravin_Kohinoor") == ("Pravin", "Kohinoor")
    assert gl.name_location("Dr K Srikanth Panjagutta") == ("Dr K Srikanth Panjagutta", "")
    assert gl.quote_no("AF - 2401-001") == "AF-2401-001"
    assert gl.division_of("Doors&Windows") == "D&W" and gl.division_of("MAP") == "MAP"
    assert gl.division_of("Sofa sets") == "Furniture"
    assert gl.vendor_name("v-33Novanthe Mobila pvt ltd po.no : AF/25-26/127") == ("Novanthe Mobila Pvt Ltd", "V33")
    assert gl.vendor_key("Arka") == gl.vendor_key("Aarka") == gl.vendor_key("AARKS")


def test_build_maps_every_sheet():
    files = [(f.filename, f.file.getvalue()) for f in _files()]
    res = gl.build(gl.load_workbooks(files))
    rec = res["records"]
    assert res["counts"]["visitors"] == 3 and res["counts"]["quotes"] == 3 and res["counts"]["sales"] == 2
    v = {x["name"]: x for x in rec["visitors"]}
    assert v["Kiran"]["date"] == "2026-01-04" and v["Kiran"]["location"] == "Gachibowli"
    assert v["Meena Rao"]["stage"] == "Won" and v["Meena Rao"]["division"] == "MAP"
    assert v["Kishore"]["stage"] == "Lost" and v["Kishore"]["date"] == "2025-09-15"   # row above's date
    q = {x["quote_no"]: x for x in rec["quotes"]}
    assert q["AF-2601-001"]["stage"] == "Won" and q["AF-2601-001"]["date"] == "2026-01-04"
    assert q["AF-2601-002"]["stage"] == "Lost" and q["AF-2601-002"]["date"] == "2026-01-01"
    assert q["AF-2601-003"]["stage"] == "Negotiation" and q["AF-2601-003"]["value"] == 120000
    s = {x["sale_no"]: x for x in rec["sales"]}
    assert s["MF 001"]["quote_id"] == q["AF-2601-001"]["id"] and s["MF 001"]["paid"] == 67546
    assert s["MF 001"]["stage"] == "Delivered" and s["MF 001"]["status"] == "PARTIAL"
    assert s["MF 002"]["paid"] == 20000 and s["MF 002"]["balance"] == 30000      # paid from balance column
    assert len(rec["customers"]) == 1 and rec["customers"][0]["phone"] == "9876500001"
    assert s["MF 001"]["customer_id"] == rec["customers"][0]["id"] == q["AF-2601-001"]["customer_id"]
    assert {a["name"] for a in rec["architects"]} == {"Ar Prakash", "Latha M"}
    pos = {p["po_no"]: p for p in rec["purchase_orders"]}
    assert pos["AF/24-25/007"]["grand_total"] == 122488 and pos["AF/24-25/007"]["status"] == "Received"
    po = pos["AF/25-26/82"]
    assert [ln["description"] for ln in po["line_items"]] == [
        "Silcol 05 kg (3 cont) (Batch : 250612AEDC02)", "Decor Matt 500 gm", "Pearl Polo"]
    assert po["grand_total"] == 8525 and po["quote_id"] == q["AF-2601-003"]["id"] and po["division"] == "MAP"
    assert pos["AF/25-26/83"]["vendor_id"] == po["vendor_id"]   # blank vendor, AHF invoice → same supplier
    inv = {i["name"]: i for i in rec["inventory"]}
    assert inv["Ficus Microcarpa"]["vendor_code"] == "V1" and inv["Ficus Microcarpa"]["status"] == "Display"
    chair = inv["Dining Chair"]                       # realigned row
    assert (chair["model_no"], chair["qty"], chair["cost"], chair["mrp"], chair["location"]) == \
           ("10115021", 4, 5500, 16500, "Godown")
    assert inv["Cimento BD 5kg"]["division"] == "MAP" and inv["Cimento BD 5kg"]["qty"] == 3
    assert res["office"] == {"gstin": "36ACDFA4831E1ZR", "name": "Madio Furniture",
                             "address": "PLOT NO 25, ROAD NO 1, KONDAPUR, HYDERABAD -500084",
                             "email": "luxury.armadio@gmail.com"}


# ── endpoints ──────────────────────────────────────────────────────────────
def test_preview_writes_nothing_and_load_needs_the_phrase():
    async def go():
        await server.db.visitors.insert_one({"id": "old", "tenant_id": "madio", "name": "Test", "date": "2026-01-01"})
        p = await server.go_live_preview(files=_files(), include_hr=False, user=ADMIN)
        assert p["counts"]["quotes"] == 3 and p["will_clear"] == {"visitors": 1}
        assert p["confirm_phrase"] == gl.CONFIRM_PHRASE and "records" not in p
        assert await server.db.quotes.count_documents({}) == 0
        with pytest.raises(HTTPException) as e:
            await server.go_live_load(files=_files(), confirm="yes", include_hr=False, apply_office=True, user=ADMIN)
        assert e.value.status_code == 400
        assert await server.db.visitors.count_documents({}) == 1
    run(go())


def test_load_archives_clears_loads_only_this_company_and_can_be_undone():
    async def go():
        db = server.db
        await db.users.insert_one({"id": "u-rooth", "tenant_id": "madio", "name": "Rooth Kumar", "role": "sales"})
        await db.visitors.insert_one({"id": "old", "tenant_id": "madio", "name": "Test", "date": "2026-01-01"})
        await db.attendance.insert_one({"id": "att", "tenant_id": "madio", "user_id": "u-rooth"})
        await db.flows.insert_one({"id": "f", "tenant_id": "madio", "name": "Mine"})
        await db.visitors.insert_one({"id": "theirs", "tenant_id": "studio", "name": "Other co", "date": "2026-01-01"})
        res = await server.go_live_load(files=_files(), confirm="delete and load", include_hr=False,
                                       apply_office=True, user=ADMIN)
        assert res["archived"] == {"visitors": 1} and res["loaded"]["visitors"] == 3
        assert await db.visitors.count_documents({"id": "old"}) == 0
        assert await db.visitors.count_documents({"id": "theirs"}) == 1           # other company untouched
        assert await db.attendance.count_documents({}) == 1 and await db.flows.count_documents({}) == 1
        v = await db.visitors.find_one({"name": "Kiran"}, {"_id": 0})
        assert v["tenant_id"] == "madio" and v["fy"] == "2025-26" and v["attend_person_id"] == "u-rooth"
        assert v["source"] == "go-live import"
        sale = await db.sales.find_one({"sale_no": "MF 001"}, {"_id": 0})
        assert sale["tenant_id"] == "madio" and sale["fy"] == "2025-26"
        office = await server._get_settings(ADMIN)
        assert office["gstin"] == "36ACDFA4831E1ZR" and office["name"] == "Madio Furniture"
        assert office["lat"] == 17.4065                                         # geofence kept

        # Another company can't see or undo it.
        assert await server.go_live_resets(user=OTHER) == []
        with pytest.raises(HTTPException):
            await server.go_live_restore(res["reset_id"], {"confirm": "RESTORE"}, user=OTHER)

        out = await server.go_live_restore(res["reset_id"], {"confirm": "restore"}, user=ADMIN)
        assert out["restored"] == {"visitors": 1}
        assert [d["id"] async for d in db.visitors.find({"tenant_id": "madio"})] == ["old"]
        assert await db.quotes.count_documents({"tenant_id": "madio"}) == 0
        assert (await server.go_live_resets(user=ADMIN))[0]["status"] == "restored"
        with pytest.raises(HTTPException):
            await server.go_live_restore(res["reset_id"], {"confirm": "RESTORE"}, user=ADMIN)
    run(go())


def test_hr_records_are_cleared_only_when_asked():
    async def go():
        await server.db.attendance.insert_one({"id": "att", "tenant_id": "madio"})
        res = await server.go_live_load(files=_files(), confirm=gl.CONFIRM_PHRASE, include_hr=True,
                                       apply_office=False, user=ADMIN)
        assert res["archived"] == {"attendance": 1} and res["office"] == {}
        assert await server.db.attendance.count_documents({}) == 0
    run(go())


def test_unrecognised_workbook_is_refused():
    async def go():
        junk = UploadFile(io.BytesIO(_book({"Notes": [("hello", "world")]})), filename="notes.xlsx")
        with pytest.raises(HTTPException) as e:
            await server.go_live_preview(files=[junk], include_hr=False, user=ADMIN)
        assert e.value.status_code == 400
        bad = UploadFile(io.BytesIO(b"not a workbook"), filename="x.xlsx")
        with pytest.raises(HTTPException):
            await server.go_live_preview(files=[bad], include_hr=False, user=ADMIN)
    run(go())


def test_starter_flows_install_once():
    async def go():
        first = await server.go_live_starter_flows(user=ADMIN)
        assert len(first["added"]) == len(server.GO_LIVE_STARTER_FLOWS) and first["skipped"] == []
        again = await server.go_live_starter_flows(user=ADMIN)
        assert again["added"] == [] and len(again["skipped"]) == len(server.GO_LIVE_STARTER_FLOWS)
        flows = await server.flows_list(user=ADMIN)
        assert all(f["tenant_id"] == "madio" for f in flows)
        assert await server.flows_list(user=OTHER) == []
    run(go())


# ── from a SharePoint folder, and the one-time automatic load ──────────────
def _fake_sharepoint(monkeypatch, files=None):
    books = files if files is not None else {f.filename: f.file.getvalue() for f in _files()}
    items = [{"name": n, "id": f"i{k}", "size": len(b), "ref": f"sharepoint:d/i{k}"}
             for k, (n, b) in enumerate(books.items())] + [
        {"name": "notes.txt", "id": "t", "size": 3, "ref": "sharepoint:d/t"}]
    by_ref = {i["ref"]: books.get(i["name"], b"txt") for i in items}
    seen = []
    monkeypatch.setattr(server.storage, "sharepoint_list_folder", lambda sub: (seen.append(sub), items)[1])
    monkeypatch.setattr(server.storage, "sharepoint_read", lambda ref: by_ref[ref])
    return seen


def test_sharepoint_folder_lists_previews_and_loads(monkeypatch):
    seen = _fake_sharepoint(monkeypatch)

    async def go():
        listing = await server.go_live_sharepoint_list(user=ADMIN)
        assert [f["name"] for f in listing["files"]] == ["AF_Sheet.xlsx", "Purchase_Order.xlsx", "MIS.xlsx"]
        p = await server.go_live_sharepoint_preview({"include_hr": False}, user=ADMIN)
        assert p["counts"]["quotes"] == 3 and p["files"] == ["AF_Sheet.xlsx", "Purchase_Order.xlsx", "MIS.xlsx"]
        with pytest.raises(HTTPException):
            await server.go_live_sharepoint_load({"confirm": "no"}, user=ADMIN)
        res = await server.go_live_sharepoint_load({"confirm": gl.CONFIRM_PHRASE}, user=ADMIN)
        assert res["loaded"]["sales"] == 2
        assert (await server.go_live_resets(user=ADMIN))[0]["source"] == "sharepoint"
    run(go())
    assert seen and all(s == "go-live" for s in seen)


def test_empty_sharepoint_folder_is_reported(monkeypatch):
    _fake_sharepoint(monkeypatch, files={})

    async def go():
        with pytest.raises(HTTPException) as e:
            await server.go_live_sharepoint_preview({}, user=ADMIN)
        assert e.value.status_code == 400 and "No .xlsx" in e.value.detail
    run(go())


def test_auto_load_runs_once_per_run_key_for_the_default_company(monkeypatch):
    _fake_sharepoint(monkeypatch)
    monkeypatch.setenv("GO_LIVE_SHAREPOINT_RUN", "golive-1")

    async def go():
        await server.db.visitors.insert_one({"id": "old", "tenant_id": server.DEFAULT_TENANT, "name": "T",
                                             "date": "2026-01-01"})
        await server.db.visitors.insert_one({"id": "theirs", "tenant_id": "studio", "name": "O", "date": "2026-01-01"})
        first = await server.go_live_auto_load()
        assert first["loaded"]["visitors"] == 3 and first["archived"] == {"visitors": 1}
        again = await server.go_live_auto_load()
        assert "already" in again["skipped"]
        assert await server.db.visitors.count_documents({"tenant_id": server.DEFAULT_TENANT}) == 3
        assert await server.db.visitors.count_documents({"tenant_id": "studio"}) == 1
        resets = await server.db.data_resets.find({}, {"_id": 0}).to_list(10)
        assert len(resets) == 1 and resets[0]["status"] == "loaded" and resets[0]["run_key"] == "golive-1"
        # It can be undone like any other load.
        admin = {**ADMIN, "tenant_id": server.DEFAULT_TENANT}
        out = await server.go_live_restore(resets[0]["id"], {"confirm": "RESTORE"}, user=admin)
        assert out["restored"] == {"visitors": 1}
    run(go())


def test_auto_load_does_nothing_without_the_setting(monkeypatch):
    monkeypatch.delenv("GO_LIVE_SHAREPOINT_RUN", raising=False)
    assert run(server.go_live_auto_load()) is None


def test_auto_load_from_packed_data(monkeypatch):
    files = [(f.filename, f.file.getvalue()) for f in _files()]
    packed = gl.pack_sheets(gl.load_workbooks(files))
    monkeypatch.setenv("GO_LIVE_SHAREPOINT_RUN", "packed-1")
    monkeypatch.setenv("GO_LIVE_DATA", packed)
    monkeypatch.setattr(server.storage, "sharepoint_list_folder",
                        lambda sub: (_ for _ in ()).throw(AssertionError("SharePoint not used")))

    async def go():
        out = await server.go_live_auto_load()
        assert out["loaded"]["quotes"] == 3 and out["loaded"]["inventory"] == 3
        v = await server.db.visitors.find_one({"name": "Kiran"}, {"_id": 0})
        assert v["date"] == "2026-01-04" and v["tenant_id"] == server.DEFAULT_TENANT   # real dates survive packing
        assert (await server.db.data_resets.find_one({}, {"_id": 0}))["source"] == "packed-auto"
        assert "already" in (await server.go_live_auto_load())["skipped"]
    run(go())


def test_bad_packed_data_is_logged_not_raised(monkeypatch):
    monkeypatch.setenv("GO_LIVE_SHAREPOINT_RUN", "bad-1")
    monkeypatch.setenv("GO_LIVE_DATA", "not-base64!!")
    assert run(server.go_live_auto_load()) is None
    assert run(server.db.data_resets.count_documents({})) == 0


def test_sharepoint_falls_back_to_files_directly_in_the_crm_folder(monkeypatch):
    books = {f.filename: f.file.getvalue() for f in _files()}
    items = [{"name": n, "id": f"i{k}", "size": len(b), "ref": f"sharepoint:d/i{k}"}
             for k, (n, b) in enumerate(books.items())]
    monkeypatch.setattr(server.storage, "sharepoint_list_folder", lambda sub: [] if sub else items)
    monkeypatch.setattr(server.storage, "sharepoint_read", lambda ref: books[items[int(ref[-1])]["name"]])

    async def go():
        listing = await server.go_live_sharepoint_list(user=ADMIN)
        assert len(listing["files"]) == 3 and not listing["folder"].endswith("go-live")
        assert (await server.go_live_sharepoint_preview({}, user=ADMIN))["counts"]["sales"] == 2
    run(go())


def test_same_sheet_name_in_two_books_and_invoice_gstin_ignored():
    other = _book({"Sheet2": [("Random", "notes", "here")],
                   "Tax Invoice": [("Furniture Sale",), ("Tax Invoice",), ("GSTIN 36BKHPC4545G1ZB",)],
                   "Stock list": [("Sl.no", "DATE", "PICTURE", "PRODUCT", "VENDOR", "MODEL NO", "QTY", "cash", "Bank",
                                   "RATE", "LP", "Taxable Value", "GST - 18%", "Purchase Cost", "Each cost",
                                   "Total cost", "Status"),
                                  (1, datetime(2024, 3, 31), None, "Ficus", "V1- ZHI RAN SHE", "Plants", 2, None, None,
                                   100, 200, 169, 31, 200, 300, 600, "godown"),
                                  (2, datetime(2026, 4, 19), None, "10211934", "Novanthe Mobila Pvt Ltd", "Dining chair",
                                   4, None, None, 50, 200, 169, 31, 200, None, 720, "Display")]})
    sheets = gl.load_workbooks([("PO.xlsx", _book(PURCHASES)), ("RP.xlsx", other)])
    assert "Sheet2 [RP.xlsx]" in sheets
    res = gl.build(sheets)
    inv = {i["name"]: i for i in res["records"]["inventory"]}
    assert inv["Cimento BD 5kg"]["division"] == "MAP"                     # the PO book's MAP list, not RP's Sheet2
    assert (inv["Ficus"]["mrp"], inv["Ficus"]["cost"], inv["Ficus"]["location"]) == (300, 100, "Godown")
    chair = inv["Dining chair"]                                          # code in PRODUCT, vendor in VENDOR
    assert (chair["model_no"], chair["vendor"], chair["mrp"]) == ("10211934", "Novanthe Mobila Pvt Ltd", 180)
    assert res["office"] == {}                                          # the invoice's GSTIN is a customer's


def test_floors_survive_and_stock_ledger_has_opening_stock():
    async def go():
        db = server.db
        await db.floors.insert_one({"id": "f1", "tenant_id": "madio", "name": "Showroom", "color": "blue"})
        res = await server.go_live_load(files=_files(), confirm=gl.CONFIRM_PHRASE, include_hr=False,
                                       apply_office=False, user=ADMIN)
        names = sorted([f["name"] async for f in db.floors.find({"tenant_id": "madio"})])
        assert names == ["Godown", "Showroom", "Warehouse"]                 # kept + one per stock location
        assert res["floors"]["added"] == ["Godown", "Warehouse"]
        moves = await db.stock_movements.find({"tenant_id": "madio"}, {"_id": 0}).to_list(50)
        inv = await db.inventory.find({"tenant_id": "madio"}, {"_id": 0}).to_list(50)
        assert {m["product_id"] for m in moves} == {i["sku"] for i in inv if i["qty"] > 0}
        summary = await server.stock_summary(user=ADMIN)
        assert summary                                                       # the ledger now has products
        # Stock screen lists newest first: furniture (dated by purchase) before MAP.
        newest = sorted(inv, key=lambda i: i["created_at"], reverse=True)
        assert newest[0]["division"] == "Furniture" and newest[-1]["division"] == "MAP"
    run(go())


def test_floors_cleared_by_an_earlier_load_come_back_from_its_archive():
    async def go():
        db = server.db
        await db.data_reset_archive.insert_one({"reset_id": "r0", "tenant_id": "madio", "collection": "floors",
                                                "doc": {"id": "f9", "tenant_id": "madio", "name": "Factory Unit",
                                                        "color": "moss"}})
        res = await server.go_live_load(files=_files(), confirm=gl.CONFIRM_PHRASE, include_hr=False,
                                       apply_office=False, user=ADMIN)
        assert res["floors"]["restored"] == 1
        assert "Factory Unit" in [f["name"] async for f in db.floors.find({"tenant_id": "madio"})]
        assert await db.floors.count_documents({"tenant_id": "studio"}) == 0
    run(go())


def _png(colour=(200, 30, 30)) -> bytes:
    from PIL import Image as PILImage
    buf = io.BytesIO()
    PILImage.new("RGB", (60, 40), colour).save(buf, "PNG")
    return buf.getvalue()


STOCK_HDR = ("Sl.no", "DATE", "PICTURE", "PRODUCT", "VENDOR", "MODEL NO", "QTY", "RATE", "Purchase Cost",
             "Each cost", "Status")


def test_stock_list_pictures_floating_over_the_sheet_reach_inventory():
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock list"
    ws["A2"] = "Stock"                     # sheet starts below a blank first row
    ws.append(STOCK_HDR)
    ws.append((1, datetime(2024, 3, 31), None, "Ficus", "V1- ZHI RAN SHE", "Plants", 1, 100, 100, 300, "godown"))
    ws.append((2, datetime(2024, 3, 31), None, "Sofa", "V2- Casa", "Sofa", 1, 100, 100, 300, "Display"))
    img = XLImage(io.BytesIO(_png()))
    ws.add_image(img, "C5")                # the Sofa row
    buf = io.BytesIO()
    wb.save(buf)
    pics = gl.sheet_images(buf.getvalue())
    assert list(pics["Stock list"]) == [5]
    inv = {i["name"]: i for i in gl.build(gl.load_workbooks([("RP.xlsx", buf.getvalue())]))["records"]["inventory"]}
    assert inv["Sofa"]["image_url"].startswith("data:image/jpeg;base64,")
    assert inv["Ficus"]["image_url"] == ""


def test_place_in_cell_pictures_are_read():
    import zipfile
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock list"
    ws.append(STOCK_HDR)
    ws.append((1, datetime(2024, 3, 31), None, "Ficus", "V1- ZHI", "Plants", 1, 100, 100, 300, "godown"))
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for n in src.namelist():
            data = src.read(n)
            if n == "xl/worksheets/sheet1.xml":    # C2 holds the picture as a rich value
                data = data.replace(b'<c r="C2"', b'<c r="C2" vm="1"', 1) if b'<c r="C2"' in data else \
                    data.replace(b'<c r="D2"', b'<c r="C2" t="e" vm="1"><v>#VALUE!</v></c><c r="D2"', 1)
            z.writestr(n, data)
        z.writestr("xl/metadata.xml",
                   '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                   '<futureMetadata name="XLRICHVALUE" count="1"><bk><extLst><ext uri="x">'
                   '<xlrd:rvb xmlns:xlrd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" i="0"/>'
                   '</ext></extLst></bk></futureMetadata>'
                   '<valueMetadata count="1"><bk><rc t="1" v="0"/></bk></valueMetadata></metadata>')
        z.writestr("xl/richData/rdrichvalue.xml",
                   '<rvData xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" count="1">'
                   '<rv s="0"><v>0</v><v>5</v></rv></rvData>')
        z.writestr("xl/richData/richValueRel.xml",
                   '<richValueRels xmlns="http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel" '
                   'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                   '<rel r:id="rId1"/></richValueRels>')
        z.writestr("xl/richData/_rels/richValueRel.xml.rels",
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="image" Target="../media/image1.png"/></Relationships>')
        z.writestr("xl/media/image1.png", _png((20, 90, 200)))
    pics = gl.sheet_images(out.getvalue())
    assert list(pics["Stock list"]) == [2]
    inv = gl.build(gl.load_workbooks([("RP.xlsx", out.getvalue())]))["records"]["inventory"]
    assert inv[0]["image_url"].startswith("data:image/jpeg")


def test_pictures_added_to_already_loaded_stock_once(monkeypatch):
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock list"
    ws.append(STOCK_HDR)
    ws.append((1, datetime(2024, 3, 31), None, "Ficus", "V1- ZHI", "Plants", 1, 100, 100, 300, "godown"))
    ws.append((2, datetime(2024, 3, 31), None, "Sofa", "V2- Casa", "Sofa", 1, 100, 100, 300, "Display"))
    ws.add_image(XLImage(io.BytesIO(_png())), "C3")
    buf = io.BytesIO()
    wb.save(buf)
    _fake_sharepoint(monkeypatch, files={"RP.xlsx": buf.getvalue()})

    async def go():
        db = server.db
        await db.inventory.insert_many([
            {"id": "a", "tenant_id": "madio", "sku": "MF-0001", "name": "Ficus", "image_url": ""},
            {"id": "b", "tenant_id": "madio", "sku": "MF-0002", "name": "Sofa", "image_url": ""},
            {"id": "c", "tenant_id": "other", "sku": "MF-0002", "name": "Sofa", "image_url": ""}])
        out = await server.go_live_auto_pictures()
        assert (out["pictures"], out["added"], out["unmatched"]) == (1, 1, 0)
        assert (await db.inventory.find_one({"id": "b"}))["image_url"].startswith("data:image/jpeg")
        assert (await db.inventory.find_one({"id": "a"}))["image_url"] == ""
        assert (await db.inventory.find_one({"id": "c"}))["image_url"] == ""          # other company untouched
        assert "skipped" in await server.go_live_auto_pictures()                       # once only
        again = await server.go_live_sharepoint_pictures({}, user=ADMIN)
        assert (again["added"], again["kept"]) == (0, 1)
    run(go())
